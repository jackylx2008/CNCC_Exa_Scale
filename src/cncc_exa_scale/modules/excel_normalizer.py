"""Read and standardize activity workbooks."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
import re
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


DATE_RANGE_PATTERN = re.compile(
    r"(?:(?P<year>\d{4})\s*年\s*)?"
    r"(?P<start_month>\d{1,2})\s*月\s*"
    r"(?P<start_day>\d{1,2})\s*(?:日|号)?"
    r"(?:\s*[-~至—–]\s*"
    r"(?:(?P<end_month>\d{1,2})\s*月\s*)?"
    r"(?P<end_day>\d{1,2})\s*(?:日|号)?)?"
)
SHEET_YEAR_PATTERN = re.compile(r"(20\d{2})")


@dataclass(frozen=True)
class ParsedDate:
    start_date: date | None
    end_date: date | None
    status: str

    @property
    def duration_days(self) -> int | None:
        if self.start_date is None or self.end_date is None:
            return None
        return (self.end_date - self.start_date).days + 1


def standardize_activity_workbook(input_path: Path, output_dir: Path) -> Path:
    """Read all sheets and write a standardized activity workbook."""
    workbook = load_workbook(input_path, data_only=True)
    rows: list[dict[str, Any]] = []

    for worksheet in workbook.worksheets:
        headers = [_clean_header(worksheet.cell(1, col).value) for col in range(1, worksheet.max_column + 1)]
        default_year = _infer_default_year(worksheet)

        for row_idx in range(2, worksheet.max_row + 1):
            values = [worksheet.cell(row_idx, col).value for col in range(1, worksheet.max_column + 1)]
            if _is_empty_row(values):
                continue

            record = {headers[index]: values[index] for index in range(len(headers))}
            meeting_time = record.get("开会时间")
            parsed = parse_meeting_date(meeting_time, default_year, workbook.epoch)

            rows.append(
                {
                    "source_sheet": worksheet.title,
                    "统计年份": default_year,
                    "source_row": row_idx,
                    "序号": record.get("序号"),
                    "活动名称": _clean_text(record.get("活动名称")),
                    "主办单位": _clean_text(record.get("主办单位")),
                    "开会时间原始值": meeting_time,
                    "开始日期": parsed.start_date.isoformat() if parsed.start_date else None,
                    "结束日期": parsed.end_date.isoformat() if parsed.end_date else None,
                    "会议天数": parsed.duration_days,
                    "规模": _normalize_number(record.get("规模")),
                    "日期解析状态": parsed.status,
                }
            )

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "standardized_activity_info.xlsx"
    csv_path = output_dir / "standardized_activity_info.csv"

    frame = pd.DataFrame(rows)
    frame.to_csv(csv_path, index=False, encoding="utf-8-sig")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name="standardized")
        worksheet = writer.book["standardized"]
        worksheet.freeze_panes = "A2"
        widths = {
            "A": 14,
            "B": 10,
            "C": 10,
            "D": 8,
            "E": 44,
            "F": 28,
            "G": 18,
            "H": 14,
            "I": 14,
            "J": 10,
            "K": 12,
            "L": 16,
        }
        for column, width in widths.items():
            worksheet.column_dimensions[column].width = width
        for cell in worksheet[1]:
            cell.style = "Headline 3"
        worksheet.auto_filter.ref = worksheet.dimensions

    return output_path


def parse_meeting_date(value: Any, default_year: int, epoch: datetime) -> ParsedDate:
    if value is None or (isinstance(value, str) and not value.strip()):
        return ParsedDate(None, None, "missing")

    if isinstance(value, datetime):
        parsed_date = value.date()
        return ParsedDate(parsed_date, parsed_date, "datetime")

    if isinstance(value, date):
        return ParsedDate(value, value, "date")

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        parsed_date = from_excel(value, epoch).date()
        if default_year:
            parsed_date = _with_year(parsed_date, default_year)
        return ParsedDate(parsed_date, parsed_date, "excel_serial")

    text = _normalize_date_text(str(value))
    matched = DATE_RANGE_PATTERN.search(text)
    if matched:
        year = int(matched.group("year") or default_year)
        start_month = int(matched.group("start_month"))
        start_day = int(matched.group("start_day"))
        end_month = int(matched.group("end_month") or start_month)
        end_day = int(matched.group("end_day") or start_day)

        try:
            start_date = date(year, start_month, start_day)
            end_year = year + 1 if (end_month, end_day) < (start_month, start_day) else year
            end_date = date(end_year, end_month, end_day)
        except ValueError:
            return ParsedDate(None, None, "invalid_text_date")

        return ParsedDate(start_date, end_date, "text_range" if end_date != start_date else "text_date")

    parsed = pd.to_datetime(text, errors="coerce")
    if pd.notna(parsed):
        parsed_date = parsed.date()
        return ParsedDate(parsed_date, parsed_date, "text_datetime")

    return ParsedDate(None, None, "unparsed")


def _infer_default_year(worksheet: Any) -> int:
    matched = SHEET_YEAR_PATTERN.search(str(worksheet.title))
    if matched:
        return int(matched.group(1))

    return datetime.now().year


def _with_year(value: date, year: int) -> date:
    try:
        return value.replace(year=year)
    except ValueError:
        return value.replace(year=year, day=28)


def _clean_header(value: Any) -> str:
    text = _clean_text(value)
    if "序号" in text:
        return "序号"
    if "活动名称" in text:
        return "活动名称"
    if "主办单位" in text:
        return "主办单位"
    if "开会时间" in text:
        return "开会时间"
    if "规模" in text:
        return "规模"
    return text


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return re.sub(r"\s+", " ", text) if text else None


def _normalize_date_text(value: str) -> str:
    return (
        value.strip()
        .replace("－", "-")
        .replace("—", "-")
        .replace("–", "-")
        .replace("~", "-")
        .replace("至", "-")
    )


def _normalize_number(value: Any) -> int | float | str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value).strip().replace(",", "")
    try:
        number = float(text)
    except ValueError:
        return text
    return int(number) if number.is_integer() else number


def _is_empty_row(values: list[Any]) -> bool:
    return all(value is None or (isinstance(value, str) and not value.strip()) for value in values)
