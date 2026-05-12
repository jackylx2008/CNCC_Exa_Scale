"""Calculate person-time statistics from standardized activity data."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, Side


QUARTER_LABELS = {1: "一季度", 2: "二季度", 3: "三季度", 4: "四季度"}
MONTH_LABELS = {
    1: "1月份",
    2: "2月份",
    3: "3月份",
    4: "4月份",
    5: "5月份",
    6: "6月份",
    7: "7月份",
    8: "8月份",
    9: "9月份",
    10: "10月份",
    11: "11月份",
    12: "12月份",
}


@dataclass(frozen=True)
class PersonTimeStatsResult:
    output_path: Path
    comparison: pd.DataFrame
    days_comparison: pd.DataFrame
    quarter_comparisons: dict[str, pd.DataFrame]
    detail: pd.DataFrame
    excluded: pd.DataFrame


def calculate_quarterly_person_times(
    standardized_path: Path,
    output_dir: Path,
    years: list[int] | None = None,
) -> PersonTimeStatsResult:
    """Calculate quarterly person-times and year-over-year comparison."""
    data = pd.read_excel(standardized_path, sheet_name="standardized")
    data["开始日期"] = pd.to_datetime(data["开始日期"], errors="coerce")
    data["结束日期"] = pd.to_datetime(data["结束日期"], errors="coerce")
    data["规模"] = pd.to_numeric(data["规模"], errors="coerce")
    if "统计年份" not in data.columns:
        data["统计年份"] = data["开始日期"].dt.year
    data["统计年份"] = pd.to_numeric(data["统计年份"], errors="coerce")

    comparison_years = years or _infer_comparison_years(data)
    detail_rows: list[dict[str, Any]] = []
    excluded_rows: list[dict[str, Any]] = []

    for _, row in data.iterrows():
        start_value = row["开始日期"]
        end_value = row["结束日期"]
        scale = row["规模"]
        year_value = row["统计年份"]

        if pd.isna(start_value) or pd.isna(end_value) or pd.isna(scale) or pd.isna(year_value):
            excluded_rows.append(_excluded_row(row, "日期、统计年份或规模为空/不可解析"))
            continue

        stat_year = int(year_value)
        if stat_year not in comparison_years:
            continue

        year_start = date(stat_year, 1, 1)
        year_end = date(stat_year, 12, 31)
        start_date = start_value.date()
        end_date = end_value.date()
        if end_date < year_start or start_date > year_end:
            continue

        counted_start = max(start_date, year_start)
        counted_end = min(end_date, year_end)
        days_by_month = _count_days_by_month(counted_start, counted_end)

        for month, days in days_by_month.items():
            quarter = (month - 1) // 3 + 1
            detail_rows.append(
                {
                    "统计年份": stat_year,
                    "季度": QUARTER_LABELS[quarter],
                    "月份": MONTH_LABELS[month],
                    "source_sheet": row.get("source_sheet"),
                    "source_row": row.get("source_row"),
                    "活动名称": row.get("活动名称"),
                    "主办单位": row.get("主办单位"),
                    "开始日期": start_date.isoformat(),
                    "结束日期": end_date.isoformat(),
                    "计入天数": days,
                    "人数": scale,
                    "人次数": scale * days,
                }
            )

    detail = pd.DataFrame(detail_rows)
    excluded = pd.DataFrame(excluded_rows)
    comparison = _build_comparison(detail, comparison_years, "人次数", "人次数")
    days_comparison = _build_comparison(detail, comparison_years, "计入天数", "计入会议天数")
    quarter_comparisons = _build_quarter_metric_comparisons(detail, comparison_years)
    sheet_summary = _build_sheet_summary(detail)

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "quarterly_person_times.xlsx"
    summary_csv_path = output_dir / "quarterly_person_times.csv"

    comparison.to_csv(summary_csv_path, index=False, encoding="utf-8-sig")
    _write_stats_workbook(
        output_path,
        comparison,
        days_comparison,
        quarter_comparisons,
        sheet_summary,
        detail,
        excluded,
        comparison_years,
    )

    return PersonTimeStatsResult(
        output_path=output_path,
        comparison=comparison,
        days_comparison=days_comparison,
        quarter_comparisons=quarter_comparisons,
        detail=detail,
        excluded=excluded,
    )


def _write_stats_workbook(
    output_path: Path,
    comparison: pd.DataFrame,
    days_comparison: pd.DataFrame,
    quarter_comparisons: dict[str, pd.DataFrame],
    sheet_summary: pd.DataFrame,
    detail: pd.DataFrame,
    excluded: pd.DataFrame,
    comparison_years: list[int],
) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        comparison.to_excel(writer, index=False, sheet_name="person_times_comparison")
        days_comparison.to_excel(writer, index=False, sheet_name="days_comparison")
        for sheet_name, quarter_comparison in quarter_comparisons.items():
            quarter_comparison.to_excel(writer, index=False, sheet_name=sheet_name)
        sheet_summary.to_excel(writer, index=False, sheet_name="sheet_quarter_summary")
        detail.to_excel(writer, index=False, sheet_name="allocation_detail")
        excluded.to_excel(writer, index=False, sheet_name="excluded")
        _format_workbook(writer.book, comparison_years)


def _infer_comparison_years(data: pd.DataFrame) -> list[int]:
    years = sorted(int(year) for year in data["统计年份"].dropna().unique())
    if len(years) >= 2:
        return years[-2:]
    return years or [date.today().year - 1, date.today().year]


def _count_days_by_month(start_date: date, end_date: date) -> dict[int, int]:
    days_by_month: dict[int, int] = {}
    current = start_date
    while current <= end_date:
        days_by_month[current.month] = days_by_month.get(current.month, 0) + 1
        current += timedelta(days=1)
    return days_by_month


def _build_comparison(
    detail: pd.DataFrame,
    years: list[int],
    value_column: str,
    metric_label: str,
) -> pd.DataFrame:
    previous_year, current_year = years[0], years[-1]
    previous_column = f"{previous_year} 年\n{metric_label}"
    current_column = f"{current_year} 年\n{metric_label}"
    rows = []
    previous_total = 0.0
    current_total = 0.0
    for quarter_label in QUARTER_LABELS.values():
        if detail.empty:
            previous_value = 0.0
            current_value = 0.0
        else:
            previous_value = _sum_value(detail, previous_year, quarter_label, value_column)
            current_value = _sum_value(detail, current_year, quarter_label, value_column)
        previous_total += previous_value
        current_total += current_value
        yoy = None if previous_value == 0 else (current_value - previous_value) / previous_value
        rows.append(
            {
                "类别": quarter_label,
                previous_column: previous_value,
                current_column: current_value,
                "增减量": current_value - previous_value,
                "同比": yoy,
            }
        )
    rows.append(
        {
            "类别": "总和",
            previous_column: previous_total,
            current_column: current_total,
            "增减量": current_total - previous_total,
            "同比": None if previous_total == 0 else (current_total - previous_total) / previous_total,
        }
    )
    return pd.DataFrame(rows)


def _build_quarter_metric_comparisons(detail: pd.DataFrame, years: list[int]) -> dict[str, pd.DataFrame]:
    comparisons = {}
    for quarter_label in QUARTER_LABELS.values():
        comparisons[f"{quarter_label}人次数"] = _build_metric_comparison_for_quarter(
            detail=detail,
            years=years,
            quarter_label=quarter_label,
            metric_label="人次数",
            value_column="人次数",
        )
        comparisons[f"{quarter_label}天数"] = _build_metric_comparison_for_quarter(
            detail=detail,
            years=years,
            quarter_label=quarter_label,
            metric_label="计入会议天数",
            value_column="计入天数",
        )
    return comparisons


def _build_metric_comparison_for_quarter(
    detail: pd.DataFrame,
    years: list[int],
    quarter_label: str,
    metric_label: str,
    value_column: str,
) -> pd.DataFrame:
    rows = []
    quarter_index = next(index for index, label in QUARTER_LABELS.items() if label == quarter_label)
    quarter_months = range((quarter_index - 1) * 3 + 1, quarter_index * 3 + 1)
    for month in quarter_months:
        month_label = MONTH_LABELS[month]
        rows.extend(
            _build_metric_rows(
                detail=detail,
                years=years,
                quarter_label=quarter_label,
                month_label=month_label,
                display_month=month_label,
                value_column=value_column,
            )
        )
    rows.extend(
        _build_metric_rows(
            detail=detail,
            years=years,
            quarter_label=quarter_label,
            month_label=None,
            display_month="季度合计",
            value_column=value_column,
        )
    )
    return pd.DataFrame(rows)


def _build_metric_rows(
    detail: pd.DataFrame,
    years: list[int],
    quarter_label: str,
    month_label: str | None,
    display_month: str,
    value_column: str,
) -> list[dict[str, float | str | None]]:
    previous_year, current_year = years[0], years[-1]
    rows: list[dict[str, float | str | None]] = []
    if detail.empty:
        previous_value = 0.0
        current_value = 0.0
    else:
        previous_value = _sum_value(detail, previous_year, quarter_label, value_column, month_label)
        current_value = _sum_value(detail, current_year, quarter_label, value_column, month_label)
    rows.append(
        {
            "月份": display_month,
            str(previous_year): previous_value,
            str(current_year): current_value,
            "增减量": current_value - previous_value,
            "同比": None if previous_value == 0 else (current_value - previous_value) / previous_value,
        }
    )
    return rows


def _sum_value(
    detail: pd.DataFrame,
    year: int,
    quarter_label: str,
    value_column: str,
    month_label: str | None = None,
) -> float:
    filtered = detail[(detail["统计年份"] == year) & (detail["季度"] == quarter_label)]
    if month_label is not None:
        filtered = filtered[filtered["月份"] == month_label]
    return float(filtered[value_column].sum())


def _build_sheet_summary(detail: pd.DataFrame) -> pd.DataFrame:
    columns = ["统计年份", "source_sheet", "季度", "活动数", "计入会议天数", "人次数"]
    if detail.empty:
        return pd.DataFrame(columns=columns)

    grouped = (
        detail.groupby(["统计年份", "source_sheet", "季度"], as_index=False)
        .agg(
            活动数=("source_row", "nunique"),
            计入会议天数=("计入天数", "sum"),
            人次数=("人次数", "sum"),
        )
        .sort_values(["统计年份", "source_sheet", "季度"])
    )
    return grouped[columns]


def _excluded_row(row: pd.Series, reason: str) -> dict[str, Any]:
    return {
        "source_sheet": row.get("source_sheet"),
        "source_row": row.get("source_row"),
        "活动名称": row.get("活动名称"),
        "开会时间原始值": row.get("开会时间原始值"),
        "开始日期": row.get("开始日期"),
        "结束日期": row.get("结束日期"),
        "规模": row.get("规模"),
        "原因": reason,
    }


def _format_workbook(workbook: Any, years: list[int]) -> None:
    widths = {
        "person_times_comparison": {"A": 18, "B": 20, "C": 20, "D": 16, "E": 12},
        "days_comparison": {"A": 18, "B": 20, "C": 20, "D": 16, "E": 12},
        "sheet_quarter_summary": {"A": 10, "B": 14, "C": 10, "D": 10, "E": 14, "F": 14},
        "allocation_detail": {
            "A": 10,
            "B": 10,
            "C": 14,
            "D": 10,
            "E": 44,
            "F": 28,
            "G": 14,
            "H": 14,
            "I": 12,
            "J": 10,
            "K": 14,
        },
        "excluded": {"A": 14, "B": 10, "C": 44, "D": 18, "E": 18, "F": 18, "G": 10, "H": 24},
    }
    titles = {
        "person_times_comparison": "1. 各季度人次数同比对比",
        "days_comparison": "2. 各季度计入会议天数同比对比",
    }
    quarter_sheet_names = [
        f"{quarter_label}{metric}"
        for quarter_label in QUARTER_LABELS.values()
        for metric in ["人次数", "天数"]
    ]
    for index, sheet_name in enumerate(quarter_sheet_names, start=3):
        metric_title = "人次数" if sheet_name.endswith("人次数") else "计入会议天数"
        quarter_label = sheet_name.removesuffix("人次数").removesuffix("天数")
        titles[sheet_name] = f"{index}. {quarter_label}{metric_title}同比对比"
        widths[sheet_name] = {"A": 14, "B": 18, "C": 18, "D": 16, "E": 12}
    for sheet_name, title in titles.items():
        worksheet = workbook[sheet_name]
        worksheet.insert_rows(1)
        worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=worksheet.max_column)
        title_cell = worksheet.cell(1, 1)
        title_cell.value = title
        title_cell.font = Font(name="宋体", size=14, bold=True)
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        worksheet.row_dimensions[1].height = 28

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        header_row = 2 if worksheet.title in titles else 1
        for cell in worksheet[header_row]:
            cell.style = "Headline 3"
        for column, width in widths.get(worksheet.title, {}).items():
            worksheet.column_dimensions[column].width = width
        _apply_borders(worksheet)

    comparison_sheet_names = ["person_times_comparison", "days_comparison", *quarter_sheet_names]
    for sheet_name in comparison_sheet_names:
        comparison = workbook[sheet_name]
        comparison.freeze_panes = "A3"
        comparison.auto_filter.ref = None
        comparison.row_dimensions[2].height = 42
        max_col = comparison.max_column
        for row in comparison.iter_rows(min_row=2, max_row=comparison.max_row, min_col=1, max_col=max_col):
            for cell in row:
                cell.font = Font(name="宋体", size=11)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for cell in comparison[2]:
            cell.font = Font(name="宋体", size=11, bold=False)
        for cell in comparison[comparison.max_row]:
            cell.font = Font(name="宋体", size=11, bold=True)
        numeric_start_col = 2
        for row in comparison.iter_rows(min_row=3, min_col=numeric_start_col, max_col=max_col - 1):
            for cell in row:
                cell.number_format = '#,##0'
        yoy_col_letter = "E" if max_col == 5 else "F"
        for cell in comparison[yoy_col_letter][2:]:
            cell.number_format = '+0.00%;-0.00%;0.00%'
            if isinstance(cell.value, (int, float)) and cell.value > 0:
                cell.font = Font(name="宋体", size=11, color="FF0000", bold=(cell.row == comparison.max_row))
        _apply_borders(comparison)


def _apply_borders(worksheet: Any) -> None:
    if worksheet.max_row < 1 or worksheet.max_column < 1:
        return
    border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    for row in worksheet.iter_rows(
        min_row=1,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=worksheet.max_column,
    ):
        for cell in row:
            cell.border = border
